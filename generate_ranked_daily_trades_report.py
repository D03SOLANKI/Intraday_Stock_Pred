import os
import sys
import pandas as pd
import numpy as np
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=40, bottom=40, left=60, right=60):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def generate_daily_ranking():
    csv_path = r'e:\stock_predictor\stock_predictor\strategy_backtest_trades_net_pnl.csv'
    df = pd.read_csv(csv_path)
    print(f"Loaded {len(df)} trades from {csv_path}...")
    
    # Aggregate day-level statistics
    daily_stats = df.groupby('trading_date').agg(
        trade_count=('trade_id', 'count'),
        total_capital=('capital_allocated', 'sum'),
        gross_pnl=('gross_pnl_inr', 'sum'),
        total_charges=('total_charges_inr', 'sum'),
        net_pnl=('net_pnl_inr', 'sum'),
        wins=('net_pnl_inr', lambda s: (s > 0).sum()),
        losses=('net_pnl_inr', lambda s: (s <= 0).sum())
    ).sort_values(by=['trade_count', 'net_pnl'], ascending=[False, False]).reset_index()
    
    daily_stats['activity_rank'] = np.arange(1, len(daily_stats) + 1)
    
    # Save day-level CSV
    daily_csv = r'e:\stock_predictor\stock_predictor\daily_trading_activity_ranked.csv'
    daily_stats.to_csv(daily_csv, index=False)
    print(f"Saved ranked daily summary to: {daily_csv} ({len(daily_stats)} active days).")
    
    # Merge activity rank back into trades dataframe
    trades_ranked = pd.merge(df, daily_stats[['trading_date', 'activity_rank', 'trade_count']], on='trading_date')
    trades_ranked = trades_ranked.sort_values(by=['activity_rank', 'trading_date', 'trade_id']).reset_index(drop=True)
    
    # Add Win/Loss flag
    trades_ranked['trade_result'] = np.where(trades_ranked['net_pnl_inr'] > 0, 'WIN', 'LOSS')
    
    # Save trade-level ranked CSV
    trades_ranked_csv = r'e:\stock_predictor\stock_predictor\daily_trades_ranked_by_activity.csv'
    trades_ranked.to_csv(trades_ranked_csv, index=False)
    print(f"Saved ranked trade-level ledger to: {trades_ranked_csv}")
    
    # Generate Word Document
    print("Building comprehensive Word document for all 209 days...")
    doc = docx.Document()
    
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)
        
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(9.5)
    normal_style.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    
    # Cover / Header
    p_t = doc.add_paragraph()
    p_t.paragraph_format.space_before = Pt(80)
    p_t.paragraph_format.space_after = Pt(8)
    r = p_t.add_run("Daily Trading Activity & Capital Deployment Audit:\nAll Backtest Days Ranked by Trades per Day")
    r.font.name = 'Arial'
    r.font.size = Pt(22)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x0F, 0x20, 0x42)
    
    p_sub = doc.add_paragraph(
        "Complete Census of All 209 Active Trading Sessions Ranked from Most Active (5 Trades) to Least Active (1 Trade), "
        "Detailing Capital Allocated, Entry, SL, TP, Quantity, Gross P&L, Deductions, Net P&L, and Win/Loss Outcomes."
    )
    p_sub.runs[0].font.italic = True
    
    doc.add_page_break()
    
    # Summary Table of Tiers
    h_tiers = doc.add_heading("1. Trading Activity Tier Distribution", level=1)
    h_tiers.runs[0].font.color.rgb = RGBColor(0x0F, 0x20, 0x42)
    
    tier_table = doc.add_table(rows=1, cols=6)
    tier_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_headers = ["Activity Tier", "Trades per Day", "Number of Days", "% of Active Days", "Total Trades Taken", "Net P&L Generated"]
    for i, th in enumerate(t_headers):
        c = tier_table.rows[0].cells[i]
        c.text = th
        c.paragraphs[0].runs[0].font.bold = True
        c.paragraphs[0].runs[0].font.size = Pt(8.5)
        c.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_background(c, "0F2042")
        set_cell_margins(c, 40, 40, 50, 50)
        
    for tc in [5, 4, 3, 2, 1]:
        sub_d = daily_stats[daily_stats['trade_count'] == tc]
        r = tier_table.add_row()
        vals = [
            f"Tier {6-tc}: Max Capacity" if tc == 5 else f"Tier {6-tc}",
            f"{tc} trades/day",
            f"{len(sub_d)} days",
            f"{(len(sub_d)/len(daily_stats))*100:.1f}%",
            f"{len(sub_d)*tc} trades",
            f"₹{sub_d['net_pnl'].sum():,.2f}"
        ]
        for c_i, val in enumerate(vals):
            c = r.cells[c_i]
            c.text = val
            c.paragraphs[0].runs[0].font.size = Pt(8.0)
            if c_i in [2, 3, 4, 5]:
                c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            set_cell_margins(c, 30, 30, 40, 40)
            
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    
    # 2. All 209 Days Ranked
    doc.add_heading("2. All 209 Trading Days Ranked by Trade Volume (With Full Trade Itemization)", level=1)
    
    # Iterate through each ranked day
    for idx_d, d_row in daily_stats.iterrows():
        d = d_row['trading_date']
        rank_num = d_row['activity_rank']
        tc = d_row['trade_count']
        tot_cap = d_row['total_capital']
        g_pnl = d_row['gross_pnl']
        chg = d_row['total_charges']
        n_pnl = d_row['net_pnl']
        wins = d_row['wins']
        losses = d_row['losses']
        
        # Day Heading
        p_dh = doc.add_paragraph()
        p_dh.paragraph_format.space_before = Pt(8)
        p_dh.paragraph_format.space_after = Pt(2)
        r_dh = p_dh.add_run(
            f"Rank #{rank_num} | Date: {d} | Total Trades: {tc} | Total Capital Deployed: ₹{tot_cap:,.2f} | "
            f"Net P&L: {'+' if n_pnl >= 0 else ''}₹{n_pnl:,.2f} ({wins}W / {losses}L)\n"
        )
        r_dh.bold = True
        r_dh.font.size = Pt(9.5)
        r_dh.font.color.rgb = RGBColor(0x0F, 0x20, 0x42)
        
        # Day trades table
        day_trades = trades_ranked[trades_ranked['trading_date'] == d]
        
        t = doc.add_table(rows=1, cols=10)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        headers = ["Symbol", "Quantity", "Entry (₹)", "SL (₹)", "TP (₹)", "Capital (₹)", "Gross (₹)", "Taxes (₹)", "Net P&L (₹)", "Result"]
        for i, h in enumerate(headers):
            cell = t.rows[0].cells[i]
            cell.text = h
            cell.paragraphs[0].runs[0].font.bold = True
            cell.paragraphs[0].runs[0].font.size = Pt(7.0)
            cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            set_cell_background(cell, "1B365D")
            set_cell_margins(cell, 20, 20, 30, 30)
            
        for _, tr in day_trades.iterrows():
            r = t.add_row()
            tr_pnl = tr['net_pnl_inr']
            vals = [
                str(tr['symbol']),
                f"{tr['shares']:,}",
                f"{tr['actual_entry_price']:.2f}",
                f"{tr['stop_loss']:.2f}",
                f"{tr['take_profit']:.2f}",
                f"{tr['capital_allocated']:,.0f}",
                f"{'+' if tr['gross_pnl_inr'] >= 0 else ''}{tr['gross_pnl_inr']:,.0f}",
                f"{tr['total_charges_inr']:.0f}",
                f"{'+' if tr_pnl >= 0 else ''}{tr_pnl:,.0f}",
                str(tr['trade_result'])
            ]
            for c_i, val in enumerate(vals):
                c = r.cells[c_i]
                c.text = val
                c.paragraphs[0].runs[0].font.size = Pt(6.5)
                if c_i in [1, 2, 3, 4, 5, 6, 7, 8]:
                    c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
                if c_i == 9:
                    c.paragraphs[0].runs[0].font.bold = True
                    if tr['trade_result'] == 'WIN':
                        set_cell_background(c, "EBF5FB")
                    else:
                        set_cell_background(c, "FDEDEC")
                set_cell_margins(c, 20, 20, 25, 25)
                
        if (idx_d + 1) % 50 == 0:
            print(f"Generated {idx_d + 1} of {len(daily_stats)} ranked days ({((idx_d + 1)/len(daily_stats))*100:.1f}%)...")
            
    doc_out = r'C:\Users\DEV SOLANKI\.gemini\antigravity\brain\1039690e-8900-4776-bdf6-ba1ff6093472\MidCap_Top_Gainer_Daily_Trade_Activity_Ranking.docx'
    doc.save(doc_out)
    print(f"Successfully saved Word report to: {doc_out}")
    
    # Copy to workspace
    dst = r'e:\stock_predictor\stock_predictor\MidCap_Top_Gainer_Daily_Trade_Activity_Ranking.docx'
    import shutil
    try:
        shutil.copyfile(doc_out, dst)
        print(f"Copied to workspace at: {dst}")
    except Exception as e:
        print(f"Copy note: {e}")

if __name__ == '__main__':
    generate_daily_ranking()
